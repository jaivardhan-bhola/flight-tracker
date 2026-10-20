#!/usr/bin/env python3
"""Extract BLR -> DEL rows (direct, via-stop, connecting) from a SpiceJet schedule .xlsx
into schedule_blr_del.csv. Re-run whenever a new schedule file is published:

    pip install openpyxl
    python extract_schedule.py path/to/Schedules_Domestic.xlsx
"""
import csv
import datetime as dt
import re
import sys

import openpyxl

VIA_IATA = {
    "MUMBAI": "BOM", "AHMEDABAD": "AMD", "KOLKATA": "CCU", "PATNA": "PAT",
    "VARANASI": "VNS", "HYDERABAD": "HYD", "PUNE": "PNQ", "JAIPUR": "JAI",
    "CHENNAI": "MAA", "GOA": "GOI", "LUCKNOW": "LKO", "GUWAHATI": "GAU",
    "BHOPAL": "BHO", "NAGPUR": "NAG", "DELHI": "DEL", "BENGALURU": "BLR",
}


def norm_flight(s: str) -> str:
    return re.sub(r"\s+", "", s or "").upper()


def pdate(s: str) -> str:
    return dt.datetime.strptime(s.strip(), "%d %b %y").date().isoformat()


def main(path: str, out: str = "schedule_blr_del.csv") -> None:
    wb = openpyxl.load_workbook(path, read_only=True)
    rows = []
    for sheet in wb.sheetnames:
        for r in wb[sheet].iter_rows(min_row=6, values_only=True):
            if not r[0] or not r[1]:
                continue
            if r[0].strip().upper() != "BENGALURU" or r[1].strip().upper() != "DELHI":
                continue
            flights = [norm_flight(x) for x in r[3].split("/")]
            via = (r[6] or "-").strip().upper()
            if len(flights) > 1:
                kind = "connecting"
            elif via not in ("-", ""):
                kind = "via_stop"
            else:
                kind = "direct"
            rows.append({
                "kind": kind,
                "flights": "|".join(flights),
                "via": "" if via == "-" else via,
                "via_iata": VIA_IATA.get(via, ""),
                "days": r[2].strip().upper(),
                "dep": r[4].strip(),
                "arr": r[5].strip(),
                "valid_from": pdate(r[7]),
                "valid_till": pdate(r[8]),
            })
    with open(out, "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)
    print(f"wrote {len(rows)} rows to {out}")


if __name__ == "__main__":
    main(sys.argv[1])
